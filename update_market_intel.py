r"""
Market Intel updater for Infra_Deal_Pipeline.xlsx

Fetches infrastructure- and energy-allocation headlines from Google News RSS (several
themed queries, merged), classifies them with the Anthropic API, and appends new
watchlist-LP mentions to the "Market Intel" sheet. Idempotent — safe to re-run.
Dedupes by article URL. Only rows dated within the last WINDOW_MONTHS are kept.

Usage:  python update_market_intel.py
        python update_market_intel.py --no-ai    (regex matching only, no API calls)

The API key is read from, in order: the ANTHROPIC_API_KEY environment variable, then
%USERPROFILE%\.anthropic\api_key.txt. That file is deliberately outside OneDrive so the
key is never cloud-synced. Without a key the script falls back to the original regex
alias matching and leaves the enrichment columns blank.

Edit WATCHLIST below to add/remove LPs. AUM figures are approximate (public sources).
"""

import argparse
import os
import re
import sys
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from typing import Literal, Optional

from dateutil.relativedelta import relativedelta

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

# ────────────────────────────────────────────────────────────────────
# Config
# ────────────────────────────────────────────────────────────────────
# Workbook lives alongside this script.
WORKBOOK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Infra_Deal_Pipeline.xlsx")
SHEET_NAME = "Market Intel"

MODEL = "claude-opus-4-8"

# Key file lives outside OneDrive so the credential is never cloud-synced.
KEY_FILE = os.path.join(os.path.expanduser("~"), ".anthropic", "api_key.txt")
KEY_PLACEHOLDER = "sk-ant-REPLACE-ME"

# Headlines per API call. See classify_headlines() — one big call exceeds the
# connection's idle tolerance and drops.
CHUNK_SIZE = 25

# Rolling window. Headlines older than this are never ingested, and existing rows
# that age past it are DELETED from the sheet on each run (see prune_and_sort).
WINDOW_MONTHS = 6

GOOGLE_NEWS_URL = "https://news.google.com/rss/search?q=%s&hl=en-US&gl=US&ceid=US:en"

# Each Google News query is an AND of three groups:
#     (fund term)  AND  (sector term — one theme below)  AND  (action word)
# Google News caps a single query at ~100 results, so the sector vocabulary is split
# into themed queries fetched separately and merged (dedup by URL). This stops
# high-volume renewables/energy coverage from crowding out a niche fund's deal.
# Full reference and held-in-reserve terms: search_terms.md.
FUND_GROUP = '(pension OR "sovereign wealth" OR "pension fund")'
ACTION_GROUP = ('(commits OR commitment OR allocates OR allocation OR invests '
                'OR mandate OR pledges)')

# Sector terms per theme. Multi-word terms MUST stay quoted — Google reads an
# unquoted space as AND, which would break the OR-group logic.
SECTOR_THEMES = {
    "general": [
        "infrastructure", "infra", '"core infrastructure"', '"real assets"',
        '"essential infrastructure"', '"infrastructure fund"', '"infrastructure debt"',
    ],
    "renewables_energy": [
        "renewable", "renewables", '"renewable energy"', '"energy transition"',
        '"clean energy"', '"new energy"', "solar", '"solar farm"', "wind",
        '"offshore wind"', '"onshore wind"', '"wind farm"', "hydro", "hydropower",
        "geothermal", "hydrogen", '"green hydrogen"', '"battery storage"',
        '"energy storage"', "biomass", '"carbon capture"', "decarbonisation",
        "decarbonization", '"net zero"', "energy", '"power generation"',
        '"power plant"', "pipeline", "pipelines", "LNG", '"natural gas"', "nuclear",
        '"small modular reactor"', "midstream",
    ],
    "digital": [
        '"data centre"', '"data center"', '"digital infrastructure"', "fibre", "fiber",
        '"fibre optic"', "broadband", '"telecom towers"', '"cell towers"', "5G",
        '"subsea cable"', '"AI infrastructure"',
    ],
    "utilities": [
        "utility", "utilities", '"water utility"', "wastewater", "desalination",
        '"district heating"', '"electricity distribution"', '"power grid"',
    ],
}


def build_query(sector_terms):
    sector_group = "(" + " OR ".join(sector_terms) + ")"
    return f"{FUND_GROUP} {sector_group} {ACTION_GROUP}"

HEADERS = [
    "Date", "Investor", "Investor Type", "AUM ($M)", "Headline", "Source", "URL",
    "Event Type", "Deal Size ($M)", "Asset Class", "Counterparty", "Confidence",
]
URL_HEADER = "URL"

EVENT_TYPES = (
    "commitment",         # new or follow-on commitment to a fund/manager
    "direct_investment",  # direct stake in an asset or project
    "acquisition",        # M&A / take-private of an asset or manager
    "allocation_target",  # stated intent to raise the infra allocation
    "mandate",            # manager selection / mandate award
    "returns_report",     # performance or results reporting
    "personnel",          # hires, departures
    "other",
)

ASSET_CLASSES = (
    "digital_infrastructure", "renewables", "energy", "transport",
    "utilities", "social_infrastructure", "diversified", "unknown",
)

# Top ~50 global pensions + SWFs by AUM, weighted toward active infra investors.
# Format: (Display Name, Type, AUM in $M approx, [aliases for headline matching])
WATCHLIST = [
    # ─── Sovereign Wealth Funds ───
    ("Norway GPFG", "Sovereign Wealth Fund", 1700000,
     ["Norges Bank Investment", "Government Pension Fund Global", "NBIM"]),
    ("China Investment Corporation", "Sovereign Wealth Fund", 1350000, ["CIC"]),
    ("Abu Dhabi Investment Authority", "Sovereign Wealth Fund", 1000000, ["ADIA"]),
    ("Kuwait Investment Authority", "Sovereign Wealth Fund", 900000, ["KIA"]),
    ("Public Investment Fund", "Sovereign Wealth Fund", 925000, ["PIF", "Saudi PIF"]),
    ("GIC", "Sovereign Wealth Fund", 800000, ["Government of Singapore Investment"]),
    ("Qatar Investment Authority", "Sovereign Wealth Fund", 475000, ["QIA"]),
    ("Temasek", "Sovereign Wealth Fund", 300000, ["Temasek Holdings"]),
    ("Mubadala", "Sovereign Wealth Fund", 280000, ["Mubadala Investment"]),
    ("Korea Investment Corporation", "Sovereign Wealth Fund", 200000, ["KIC"]),
    ("Future Fund", "Sovereign Wealth Fund", 170000, ["Australian Future Fund"]),
    ("New Zealand Superannuation Fund", "Sovereign Wealth Fund", 55000, ["NZ Super Fund", "NZSF"]),
    ("Khazanah Nasional", "Sovereign Wealth Fund", 40000, ["Khazanah"]),
    ("Alaska Permanent Fund", "Sovereign Wealth Fund", 80000, ["APFC"]),

    # ─── Canadian pensions (major infra players) ───
    ("CPP Investments", "Pension Fund", 600000, ["CPPIB", "Canada Pension Plan"]),
    ("CDPQ", "Pension Fund", 330000, ["Caisse de depot", "Caisse de dépôt"]),
    ("Ontario Teachers", "Pension Fund", 250000, ["OTPP", "Ontario Teachers Pension"]),
    ("PSP Investments", "Pension Fund", 265000, ["Public Sector Pension"]),
    ("BCI", "Pension Fund", 250000, ["British Columbia Investment"]),
    ("AIMCo", "Pension Fund", 150000, ["Alberta Investment Management"]),
    ("OMERS", "Pension Fund", 130000, ["Ontario Municipal Employees"]),
    ("HOOPP", "Pension Fund", 110000, ["Healthcare of Ontario Pension"]),
    ("IMCO", "Pension Fund", 80000, ["Investment Management Corporation of Ontario"]),

    # ─── US pensions ───
    ("CalPERS", "Pension Fund", 500000, ["California Public Employees"]),
    ("CalSTRS", "Pension Fund", 330000, ["California State Teachers"]),
    ("New York State Common Retirement", "Pension Fund", 260000, ["NYSCRF", "NY State Common"]),
    ("NYC Retirement Systems", "Pension Fund", 290000, ["New York City Retirement"]),
    ("Florida SBA", "Pension Fund", 260000, ["Florida State Board of Administration"]),
    ("Texas TRS", "Pension Fund", 200000, ["Teacher Retirement System of Texas"]),
    ("Washington State Investment Board", "Pension Fund", 175000, ["WSIB"]),
    ("Wisconsin Investment Board", "Pension Fund", 155000, ["SWIB"]),
    ("Oregon PERS", "Pension Fund", 95000, ["Oregon Public Employees"]),
    ("Pennsylvania PSERS", "Pension Fund", 75000, ["Public School Employees' Retirement System"]),
    ("Virginia Retirement System", "Pension Fund", 100000, ["VRS"]),

    # ─── European pensions ───
    ("ABP", "Pension Fund", 550000, ["Stichting Pensioenfonds ABP"]),
    ("PFZW", "Pension Fund", 300000, ["Pensioenfonds Zorg en Welzijn"]),
    ("PGGM", "Pension Fund", 250000, []),
    ("ATP", "Pension Fund", 100000, ["Arbejdsmarkedets Tillaegspension"]),
    ("Alecta", "Pension Fund", 110000, []),
    ("AP2", "Pension Fund", 50000, ["Andra AP-fonden", "Second Swedish National"]),
    ("USS", "Pension Fund", 90000, ["Universities Superannuation Scheme"]),
    ("Railpen", "Pension Fund", 45000, ["Railways Pension"]),
    ("BT Pension Scheme", "Pension Fund", 55000, ["BTPS"]),

    # ─── Asia-Pacific pensions ───
    ("GPIF", "Pension Fund", 1500000, ["Government Pension Investment Fund"]),
    ("National Pension Service", "Pension Fund", 800000, ["NPS Korea", "Korean National Pension"]),
    ("AustralianSuper", "Pension Fund", 240000, []),
    ("Aware Super", "Pension Fund", 110000, []),
    ("UniSuper", "Pension Fund", 95000, []),
    ("Cbus", "Pension Fund", 65000, []),
    ("HESTA", "Pension Fund", 55000, []),

    # ─── Other ───
    ("Federal Retirement Thrift", "Pension Fund", 800000, ["Thrift Savings Plan", "TSP"]),
]

LOOKUP = {name: (name, itype, aum) for name, itype, aum, _ in WATCHLIST}


def load_api_key():
    """Resolve the API key from the environment, then the key file.

    Returns (key, source) or (None, reason). The key itself is never printed.
    """
    for var in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN"):
        val = os.environ.get(var, "").strip()
        if val:
            return val, f"${var}"

    if not os.path.exists(KEY_FILE):
        return None, f"no key file at {KEY_FILE}"

    try:
        with open(KEY_FILE, encoding="utf-8") as f:
            lines = [ln.strip() for ln in f]
    except OSError as e:
        return None, f"could not read {KEY_FILE}: {e}"

    for ln in lines:
        if not ln or ln.startswith("#"):
            continue
        if ln == KEY_PLACEHOLDER:
            return None, f"{KEY_FILE} still contains the placeholder — paste your real key in"
        return ln, KEY_FILE

    return None, f"{KEY_FILE} is empty"


def _fetch_one(query):
    """Fetch and parse one Google News RSS query into a list of headline dicts."""
    url = GOOGLE_NEWS_URL % urllib.parse.quote(query)
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read()
    root = ET.fromstring(data)
    items = []
    for item in root.iter("item"):
        title = (item.findtext("title") or "").strip()
        link = (item.findtext("link") or "").strip()
        pub = item.findtext("pubDate") or ""
        src_el = item.find("source")
        source = (src_el.text or "").strip() if src_el is not None else ""
        try:
            dt = parsedate_to_datetime(pub).date()
        except Exception:
            dt = None
        items.append({"title": title, "link": link, "date": dt, "source": source})
    return items


def fetch_headlines():
    """Run each themed query, merge results, dedupe by URL.

    A headline can match more than one theme (e.g. a data-centre-for-renewables
    deal); the first theme to see a URL wins, so each is classified only once.
    Partial failures are tolerated — only if every theme errors do we raise.
    """
    merged = {}
    errors = 0
    for theme, terms in SECTOR_THEMES.items():
        try:
            items = _fetch_one(build_query(terms))
        except Exception as e:
            print(f"  [{theme}] fetch failed ({type(e).__name__}) — skipped")
            errors += 1
            continue
        new = 0
        for it in items:
            if it["link"] and it["link"] not in merged:
                merged[it["link"]] = it
                new += 1
        print(f"  [{theme}] {len(items)} raw, {new} new after merge")
    if errors == len(SECTOR_THEMES):
        raise RuntimeError("all themed queries failed")
    return list(merged.values())


def match_watchlist(headline):
    """Regex fallback. Return (name, type, aum) if any watchlist LP is mentioned."""
    for name, itype, aum, aliases in WATCHLIST:
        for c in [name] + aliases:
            # Short all-caps acronyms: case-sensitive match to avoid false positives.
            if c.isupper() and len(c) <= 5:
                if re.search(r"\b" + re.escape(c) + r"\b", headline):
                    return name, itype, aum
            else:
                if re.search(r"\b" + re.escape(c) + r"\b", headline, re.IGNORECASE):
                    return name, itype, aum
    return None


# ────────────────────────────────────────────────────────────────────
# Claude classification
# ────────────────────────────────────────────────────────────────────
def _build_schema_models():
    """Imported lazily so --no-ai works without the anthropic package."""
    from pydantic import BaseModel, Field

    class Article(BaseModel):
        index: int = Field(description="The headline's index from the input list.")
        investor: str = Field(
            description="Exact watchlist display name, or empty string if no "
                        "watchlist institution is the subject of the story."
        )
        event_type: Literal[EVENT_TYPES]  # type: ignore[valid-type]
        deal_size_musd: Optional[float] = Field(
            description="Deal or commitment size in $ millions, or null if not stated."
        )
        asset_class: Literal[ASSET_CLASSES]  # type: ignore[valid-type]
        counterparty: str = Field(
            description="Manager, fund, or asset on the other side of the deal; "
                        "empty string if none."
        )
        confidence: Literal["high", "medium", "low"]  # type: ignore[valid-type]

    class Batch(BaseModel):
        results: list[Article]

    return Batch


SYSTEM_PROMPT = """You classify infrastructure investment news headlines for an LP watchlist.

For each headline you are given, decide whether one of the watchlist institutions below
is the SUBJECT of the story — the institution doing the investing, allocating, or reporting.

Rules:
- Match indirect references. "The Canadian pension giant", "Norway's oil fund", and
  "the $1.7trn sovereign fund" all resolve to watchlist names when context makes the
  institution unambiguous.
- If the institution is only mentioned in passing, or the subject is not on the watchlist,
  return an empty string for `investor`.
- `investor` must be an EXACT display name from the watchlist below — never an alias,
  ticker, or paraphrase.
- Classify what the story actually reports. A performance or results story is
  `returns_report` even when it mentions infrastructure gains. Only use `commitment`,
  `direct_investment`, or `acquisition` when capital is actually being deployed.
- `deal_size_musd` is in $ millions: "$2bn" -> 2000, "$300m" -> 300. Null if not stated.
  Do not infer or estimate a size that is not in the headline.
- Set `confidence` to reflect how sure you are of the investor + event_type pairing.
  A headline that names the institution and the action explicitly is `high`. An indirect
  reference, or an event type you inferred from thin wording, is `medium` or `low`.

Watchlist display names:
{watchlist}
"""


def classify_headlines(items, api_key, model=MODEL, chunk_size=CHUNK_SIZE):
    """Classify headlines in chunks. Returns ({index: Article}, {failed indices}).

    Chunked deliberately: a single ~86-headline call runs >3 minutes with no bytes
    on the wire and the connection gets dropped (APIConnectionError). ~25 headlines
    per call keeps each request around 35s. Chunking also contains failures — a bad
    chunk costs only its own headlines, which are left for the next run.
    """
    import anthropic

    Batch = _build_schema_models()
    client = anthropic.Anthropic(api_key=api_key, timeout=180.0, max_retries=2)
    names = "\n".join(f"- {name}" for name, _, _, _ in WATCHLIST)

    classified = {}
    failed = set()
    tok_in = tok_out = 0

    for start in range(0, len(items), chunk_size):
        idxs = list(range(start, min(start + chunk_size, len(items))))
        listing = "\n".join(f"[{i}] {items[i]['title']}" for i in idxs)
        label = f"{idxs[0]}-{idxs[-1]}"

        try:
            response = client.messages.parse(
                model=model,
                max_tokens=16000,
                thinking={"type": "adaptive"},
                system=SYSTEM_PROMPT.format(watchlist=names),
                messages=[{
                    "role": "user",
                    "content": f"Classify each headline. Return exactly one result per "
                               f"headline, preserving the index.\n\n{listing}",
                }],
                output_format=Batch,
            )
        except Exception as e:
            print(f"    [{label}] failed ({type(e).__name__}) — left for next run")
            failed.update(idxs)
            continue

        valid = set(idxs)
        got = {a.index: a for a in response.parsed_output.results if a.index in valid}
        classified.update(got)
        # Any headline the model silently dropped is a failure, not a "no match".
        failed.update(valid - got.keys())

        tok_in += response.usage.input_tokens
        tok_out += response.usage.output_tokens
        print(f"    [{label}] {len(got)}/{len(idxs)} classified")

    print(f"  Claude: {tok_in} in / {tok_out} out tokens across "
          f"{(len(items) + chunk_size - 1) // chunk_size} call(s)")
    return classified, failed


# ────────────────────────────────────────────────────────────────────
# Sheet handling
# ────────────────────────────────────────────────────────────────────
NAVY = "1B2A4A"


def _style_header(ws):
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [12, 28, 22, 14, 70, 22, 50, 18, 15, 22, 28, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def ensure_market_intel_sheet(wb):
    """Create the sheet, or extend an existing 7-column sheet with the new columns."""
    if SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_NAME)
        ws.sheet_properties.tabColor = "C00000"
        ws.append(HEADERS)
        _style_header(ws)
        ws.freeze_panes = "A2"
        return ws

    ws = wb[SHEET_NAME]
    existing = [c.value for c in ws[1]]
    if existing[:len(HEADERS)] != HEADERS:
        for i, h in enumerate(HEADERS, 1):
            ws.cell(row=1, column=i, value=h)
        _style_header(ws)
        print(f"  Extended '{SHEET_NAME}' header to {len(HEADERS)} columns "
              f"(existing rows keep blank enrichment cells)")
    return ws


def url_column(ws):
    for i, cell in enumerate(ws[1], 1):
        if cell.value == URL_HEADER:
            return i
    raise RuntimeError(f"No '{URL_HEADER}' column found in '{SHEET_NAME}'")


def existing_urls(ws, url_col):
    urls = set()
    for row in ws.iter_rows(min_row=2, min_col=url_col, max_col=url_col, values_only=True):
        if row[0]:
            urls.add(row[0])
    return urls


def _as_date(v):
    """Normalise a cell value to a date. openpyxl hands back datetime; RSS gives date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _format_row(ws, r):
    ws.cell(row=r, column=1).number_format = "YYYY-MM-DD"
    ws.cell(row=r, column=4).number_format = '$#,##0;($#,##0);"-"'
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=9).number_format = '$#,##0;($#,##0);"-"'


def prune_and_sort(ws, cutoff):
    """Delete rows older than `cutoff`, then rewrite the sheet newest-first.

    Destructive by design — the 6-month window is a hard rule (WINDOW_MONTHS).
    Rows with an unreadable date are KEPT and sorted to the bottom: we can't prove
    they're stale, so we don't delete them.

    Returns (pruned_rows, kept_count).
    """
    rows = [list(r) for r in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True)]
    rows = [r for r in rows if any(c is not None for c in r)]  # ignore blank rows

    kept, pruned = [], []
    for r in rows:
        d = _as_date(r[0])
        if d is not None and d < cutoff:
            pruned.append(r)
        else:
            kept.append(r)

    # Newest first; undated rows last.
    kept.sort(key=lambda r: (_as_date(r[0]) is None, -(_as_date(r[0]) or date.min).toordinal()))

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in kept:
        ws.append(r)
    for i in range(len(kept)):
        _format_row(ws, 2 + i)

    # Header dropdowns so the window can be narrowed further without code.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"
    return pruned, len(kept)


def check_file_writable(path):
    """Return True if we can open the file for writing (i.e. not locked by Excel)."""
    try:
        with open(path, "r+b"):
            return True
    except PermissionError:
        return False


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--no-ai", action="store_true",
                    help="Skip Claude classification; use regex alias matching only.")
    ap.add_argument("--model", default=MODEL, help=f"Claude model (default: {MODEL})")
    args = ap.parse_args()

    if not os.path.exists(WORKBOOK):
        print(f"ERROR: workbook not found at {WORKBOOK}")
        sys.exit(1)

    if not check_file_writable(WORKBOOK):
        print("=" * 60)
        print("  The workbook is currently open in Excel (or locked).")
        print("  Please close 'Infra_Deal_Pipeline.xlsx' and try again.")
        print("=" * 60)
        sys.exit(1)

    print("Fetching Google News RSS...")
    try:
        headlines = fetch_headlines()
    except Exception as e:
        print(f"ERROR fetching feed: {e}")
        sys.exit(1)
    print(f"  {len(headlines)} headlines retrieved")

    wb = load_workbook(WORKBOOK)
    ws = ensure_market_intel_sheet(wb)
    url_col = url_column(ws)
    seen = existing_urls(ws, url_col)

    cutoff = date.today() - relativedelta(months=WINDOW_MONTHS)

    # Filter BEFORE classifying — never pay to classify a headline we'd discard.
    candidates = []
    n_dupe = n_old = n_undated = 0
    for h in headlines:
        if not h["link"] or h["link"] in seen:
            n_dupe += 1
        elif h["date"] is None:
            n_undated += 1          # can't prove it's in-window; leave it out
        elif h["date"] < cutoff:
            n_old += 1
        else:
            candidates.append(h)
    print(f"  {len(candidates)} to classify — skipped {n_dupe} already in sheet, "
          f"{n_old} older than {cutoff}, {n_undated} undated")

    # Decide whether to classify with Claude.
    use_ai = not args.no_ai
    api_key = None
    if use_ai and candidates:
        api_key, source = load_api_key()
        if api_key:
            print(f"  API key loaded from {source}")
        else:
            print(f"  No API key ({source}) — falling back to regex matching")
            use_ai = False

    classified, failed = {}, set()
    if use_ai and candidates:
        print(f"Classifying {len(candidates)} headlines with {args.model} "
              f"(chunks of {CHUNK_SIZE})...")
        classified, failed = classify_headlines(candidates, api_key, model=args.model)
        if not classified and failed:
            print("  Every chunk failed — nothing written. Re-run, or use --no-ai "
                  "to write unenriched rows anyway.")
            sys.exit(1)

    new_rows = []
    unknown_names = set()
    skipped_failed = 0
    for i, h in enumerate(candidates):
        if use_ai:
            # A headline whose chunk failed must NOT be written unenriched: dedupe is
            # by URL, so an unenriched row is permanent. Skip it and let the next run
            # pick it up.
            if i in failed:
                skipped_failed += 1
                continue
            a = classified.get(i)
            if not a or not a.investor:
                continue
            hit = LOOKUP.get(a.investor)
            if not hit:
                unknown_names.add(a.investor)
                continue
            name, itype, aum = hit
            extra = [a.event_type, a.deal_size_musd, a.asset_class,
                     a.counterparty or None, a.confidence]
        else:
            m = match_watchlist(h["title"])
            if not m:
                continue
            name, itype, aum = m
            extra = [None] * 5

        row = [h["date"], name, itype, aum, h["title"], h["source"], h["link"]] + extra
        ws.append(row)
        seen.add(h["link"])
        new_rows.append(row)

    # Enforce the rolling window and re-sort newest-first. Destructive: rows that
    # have aged past the cutoff are deleted here.
    pruned, kept = prune_and_sort(ws, cutoff)

    try:
        wb.save(WORKBOOK)
    except PermissionError:
        print("=" * 60)
        print("  Couldn't save — the workbook was opened during the run.")
        print("  Close 'Infra_Deal_Pipeline.xlsx' and re-run the script.")
        print("=" * 60)
        sys.exit(1)

    if unknown_names:
        print(f"  Skipped {len(unknown_names)} unrecognised investor name(s) from Claude: "
              f"{', '.join(sorted(unknown_names))}")
    if skipped_failed:
        print(f"  Left {skipped_failed} headline(s) unwritten (classification failed) "
              f"— re-run to pick them up")
    if pruned:
        print(f"  Deleted {len(pruned)} row(s) older than {cutoff}:")
        for r in sorted(pruned, key=lambda x: (_as_date(x[0]) or date.min)):
            print(f"    {str(_as_date(r[0]))}  {(r[1] or '?')[:34]}")
    print(f"  {len(new_rows)} new announcement(s) added to '{SHEET_NAME}' "
          f"— {kept} row(s) now in the {WINDOW_MONTHS}-month window, newest first")
    if new_rows:
        print("\nNew entries:")
        for r in new_rows[:10]:
            tag = f" [{r[7]}]" if r[7] else ""
            print(f"  [{r[0]}] {r[1]}{tag}: {r[4][:80]}")
        if len(new_rows) > 10:
            print(f"  ... and {len(new_rows) - 10} more")


if __name__ == "__main__":
    main()
